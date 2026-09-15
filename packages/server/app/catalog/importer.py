"""公开电站目录导入：解析 WRI CSV / GEM xlsx，去重后写库。docs/04 §七

脚本 scripts/import_catalog.py 与定时任务 sync_catalog 都走这里。
"""

from __future__ import annotations

import csv
import hashlib
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import CatalogPlant
from app.services.catalog import haversine_km

DEDUP_KM = 1.5
MIN_CAPACITY_MW = 1.0
_JUNK = re.compile(r"[{}\[\]|<>]")  # WRI 源里偶有 "Qili}" 这种脏字符


def clean_name(s: str) -> str:
    return re.sub(r"\s+", " ", _JUNK.sub("", s)).strip()


def first_name(s: str) -> str:
    """GEM 的本地名会把多期名字用逗号拼在一起（常常还是重复的），取第一个不重复的。"""
    parts = [clean_name(p) for p in re.split(r"[,，;；/|]", s)]
    parts = [p for p in parts if p]
    return parts[0] if parts else ""


@dataclass
class Row:
    source: str
    source_id: str
    name: str
    name_local: str | None
    type: str  # solar | wind
    capacity_mw: float
    lat: float
    lon: float
    province: str | None = None
    city: str | None = None
    district: str | None = None
    owner: str | None = None
    year: int | None = None
    status: str = "operating"
    provenance: dict = field(default_factory=dict)

    @property
    def id(self) -> str:
        return f"{self.source}:{self.source_id}"[:24]


# ────────────────────────────── WRI ──────────────────────────────

_WRI_FUEL = {"Solar": "solar", "Wind": "wind"}


def read_wri(path: Path, country: str | None) -> list[Row]:
    out: list[Row] = []
    with path.open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if country and r["country"] != country:
                continue
            t = _WRI_FUEL.get(r["primary_fuel"])
            if not t or not r["capacity_mw"] or not r["latitude"]:
                continue
            cap = float(r["capacity_mw"])
            if cap < MIN_CAPACITY_MW:
                continue
            year = r.get("commissioning_year") or None
            out.append(
                Row(
                    source="wri",
                    source_id=r["gppd_idnr"],
                    name=clean_name(r["name"]),
                    name_local=None,
                    type=t,
                    capacity_mw=cap,
                    lat=float(r["latitude"]),
                    lon=float(r["longitude"]),
                    owner=(r.get("owner") or "").strip() or None,
                    year=int(float(year)) if year else None,
                )
            )
    return out


# ────────────────────────────── GEM ──────────────────────────────


# GEM 的省份是英文，映射成中文；市县暂留英文，配了腾讯 key 后由回填任务逐步换成中文
PROVINCE_ZH = {
    "Beijing": "北京市",
    "Tianjin": "天津市",
    "Hebei": "河北省",
    "Shanxi": "山西省",
    "Inner Mongolia": "内蒙古自治区",
    "Liaoning": "辽宁省",
    "Jilin": "吉林省",
    "Heilongjiang": "黑龙江省",
    "Shanghai": "上海市",
    "Jiangsu": "江苏省",
    "Zhejiang": "浙江省",
    "Anhui": "安徽省",
    "Fujian": "福建省",
    "Jiangxi": "江西省",
    "Shandong": "山东省",
    "Henan": "河南省",
    "Hubei": "湖北省",
    "Hunan": "湖南省",
    "Guangdong": "广东省",
    "Guangxi": "广西壮族自治区",
    "Hainan": "海南省",
    "Chongqing": "重庆市",
    "Sichuan": "四川省",
    "Guizhou": "贵州省",
    "Yunnan": "云南省",
    "Tibet": "西藏自治区",
    "Xizang": "西藏自治区",
    "Shaanxi": "陕西省",
    "Gansu": "甘肃省",
    "Qinghai": "青海省",
    "Ningxia": "宁夏回族自治区",
    "Xinjiang": "新疆维吾尔自治区",
    "Hong Kong": "香港",
    "Macau": "澳门",
    "Taiwan": "台湾",
}


def province_zh(v: str | None) -> str | None:
    if not v:
        return None
    v = str(v).strip()
    return PROVINCE_ZH.get(v, v)


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h).lower())


def _col(headers: list[str], *candidates: str) -> int | None:
    """按规范化后的列名找列；候选按优先级，支持前缀匹配。"""
    normed = [_norm(h) for h in headers]
    for c in candidates:
        cn = _norm(c)
        for i, h in enumerate(normed):
            if h == cn:
                return i
    for c in candidates:
        cn = _norm(c)
        for i, h in enumerate(normed):
            if h.startswith(cn):
                return i
    return None


def read_gem(path: Path, country: str | None) -> list[Row]:
    import openpyxl

    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # 数据表通常叫 Data；找第一张含 Latitude 列的表
    ws = None
    for sheet in wb.worksheets:
        head = [c for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        if _col([str(h) for h in head if h is not None], "Latitude") is not None:
            ws = sheet
            break
    if ws is None:
        raise SystemExit("没找到含 Latitude 列的数据表")
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows)]
    ix = {
        "country": _col(headers, "Country/Area", "Country"),
        "name": _col(headers, "Project Name", "Plant Name", "Name"),
        "local": _col(
            headers, "Project Name in Local Language / Script", "Local Name", "Chinese Name"
        ),
        "cap": _col(headers, "Capacity (MW)", "Capacity"),
        "status": _col(headers, "Status"),
        "lat": _col(headers, "Latitude"),
        "lon": _col(headers, "Longitude"),
        "owner": _col(headers, "Owner Name in Local Language / Script", "Owner", "Operator"),
        "owner_en": _col(headers, "Owner", "Operator"),
        "year": _col(headers, "Start year", "Start Year", "Year"),
        "prov": _col(headers, "State/Province", "Subnational unit (province, state)", "State"),
        "city": _col(headers, "Major area (prefecture, district)", "City"),
        "district": _col(headers, "Local area (taluk, county)", "Local area"),
        "loc_id": _col(headers, "GEM location ID", "GEM location"),
        "rating": _col(headers, "Capacity Rating"),
        "phase_name": _col(headers, "Phase Name"),
        "wiki": _col(headers, "Wiki URL"),
        "accuracy": _col(headers, "Location accuracy"),
        "phase_id": _col(headers, "GEM phase ID", "GEM unit/phase ID", "GEM unit ID"),
        "installation": _col(headers, "Installation Type"),  # 风电表：Onshore / Offshore …
    }
    missing = [k for k in ("name", "cap", "lat", "lon", "status") if ix[k] is None]
    if missing:
        raise SystemExit(f"GEM 表缺少列 {missing}，表头：{headers}")
    t = "solar" if "solar" in path.name.lower() else "wind" if "wind" in path.name.lower() else None
    tech = _col(headers, "Technology Type", "Type")
    country_alias = {"CHN": "china"}

    def cell(r, k):
        i = ix[k]
        return r[i] if i is not None and i < len(r) else None

    grouped: dict[str, Row] = {}
    phase_seen: dict[str, tuple] = {}
    for r in rows:
        if country:
            c = str(cell(r, "country") or "").lower()
            if country_alias.get(country, country.lower()) not in c:
                continue
        status = str(cell(r, "status") or "").lower()
        if status not in ("operating",):
            continue
        cap, lat, lon = cell(r, "cap"), cell(r, "lat"), cell(r, "lon")
        if cap is None or lat is None or lon is None:
            continue
        try:
            cap, lat, lon = float(cap), float(lat), float(lon)
        except (TypeError, ValueError):
            continue
        if cap < MIN_CAPACITY_MW:
            continue
        kind = t
        if kind is None and tech is not None:
            kind = "wind" if "wind" in str(r[tech] or "").lower() else "solar"
        if kind is None:
            continue
        loc = str(cell(r, "loc_id") or cell(r, "phase_id") or f"{lat:.4f},{lon:.4f}")
        year = cell(r, "year")
        try:
            year = int(float(year)) if year not in (None, "") else None
        except (TypeError, ValueError):
            year = None
        key = f"{kind}:{loc}"
        pid = str(cell(r, "phase_id") or "")
        rating_raw = str(cell(r, "rating") or "unknown")
        rating = (
            "ac"
            if rating_raw.lower() == "mwac"
            else "dc"
            if rating_raw.lower() in ("mwp/dc", "mwdc", "mwp")
            else "unknown"
        )
        signature = (loc, kind, cap, lat, lon, rating)
        if pid and pid in phase_seen:
            if phase_seen[pid] != signature:
                raise ValueError(f"同一分期 {pid} 存在冲突，拒绝覆盖目录")
            continue
        if pid:
            phase_seen[pid] = signature
        phase = dict(
            id=pid,
            name=first_name(str(cell(r, "local") or cell(r, "name") or "")),
            phase_name=str(cell(r, "phase_name") or ""),
            capacity_kw=cap * 1000,
            capacity_rating=rating,
            capacity_rating_raw=rating_raw,
            latitude=lat,
            longitude=lon,
            technology=str(r[tech] or "") if tech is not None else "",
            owner=str(cell(r, "owner") or cell(r, "owner_en") or ""),
            status=status,
            # 目录风电按投运年份与海上/陆上选机型档和气象格点，见 prediction_basis
            start_year=year,
            installation_type=str(cell(r, "installation") or ""),
        )
        if key in grouped:
            grouped[key].capacity_mw += cap
            grouped[key].provenance["phases"].append(phase)
            continue
        grouped[key] = Row(
            source="gem",
            source_id=loc,
            provenance=dict(
                source_file=path.name,
                source_sha256=digest,
                location_id=loc,
                wiki_url=str(cell(r, "wiki") or ""),
                location_accuracy=str(cell(r, "accuracy") or ""),
                phases=[phase],
            ),
            name=first_name(str(cell(r, "name") or "")),
            name_local=(first_name(str(cell(r, "local"))) or None) if cell(r, "local") else None,
            type=kind,
            capacity_mw=cap,
            lat=lat,
            lon=lon,
            province=province_zh(cell(r, "prov")),
            city=(str(cell(r, "city")).strip() or None) if cell(r, "city") else None,
            district=(str(cell(r, "district")).strip() or None) if cell(r, "district") else None,
            owner=(first_name(str(cell(r, "owner") or cell(r, "owner_en") or "")) or None),
            year=year,
        )
    for item in grouped.values():
        phases = item.provenance["phases"]
        if len(phases) > 1:
            # 不能继续把首期名称当作全部容量的名称；各期完整名称保存在溯源中。
            from app.catalog.quality import merged_site_name

            region = item.district or item.city or item.province or "公开"
            item.name_local = merged_site_name(phases, item.type, region)
    return list(grouped.values())


# ────────────────────────────── 写库 ──────────────────────────────


@dataclass
class ImportResult:
    added: int = 0
    updated: int = 0
    replaced: int = 0
    seen_ids: set[str] | None = None


async def upsert(db: AsyncSession, rows: list[Row]) -> ImportResult:
    """写库但不 commit，调用方决定事务边界。"""
    added = updated = replaced = 0
    seen: set[str] = set()
    existing = (await db.execute(select(CatalogPlant))).scalars().all()
    by_id = {p.id: p for p in existing}
    # 粗网格索引做同址查找
    grid: dict[tuple[str, int, int], list[CatalogPlant]] = {}
    for p in existing:
        grid.setdefault((p.type, int(p.latitude * 20), int(p.longitude * 20)), []).append(p)

    def near_dupes(r: Row) -> list[CatalogPlant]:
        out = []
        gy, gx = int(r.lat * 20), int(r.lon * 20)
        for dy in (-1, 0, 1):
            for dx in (-1, 0, 1):
                for p in grid.get((r.type, gy + dy, gx + dx), []):
                    if (
                        p.id != r.id
                        and haversine_km(r.lat, r.lon, p.latitude, p.longitude) < DEDUP_KM
                    ):
                        out.append(p)
        return out

    for r in rows:
        seen.add(r.id)
        dupes = near_dupes(r)
        if dupes:
            if r.source == "wri" and any(p.source == "gem" for p in dupes):
                replaced += 1
                continue  # GEM 已有更好的记录
            if r.source == "gem":
                for p in dupes:
                    if p.source == "wri":
                        await db.delete(p)
                        replaced += 1
        p = by_id.get(r.id)
        fields = dict(
            source=r.source,
            source_id=r.source_id,
            name=r.name[:128],
            name_local=(r.name_local or None) and r.name_local[:128],
            type=r.type,
            capacity_kw=r.capacity_mw * 1000.0,
            provenance=r.provenance or None,
            latitude=r.lat,
            longitude=r.lon,
            province=r.province,
            city=r.city,
            district=r.district,
            owner_name=(r.owner or None) and r.owner[:128],
            commissioning_year=r.year,
            status=r.status,
            updated_at=datetime.now(UTC).replace(tzinfo=None),
        )
        if p:
            for k, v in fields.items():
                setattr(p, k, v)
            updated += 1
        else:
            p = CatalogPlant(id=r.id, **fields)
            db.add(p)
            by_id[r.id] = p
            grid.setdefault((r.type, int(r.lat * 20), int(r.lon * 20)), []).append(p)
            added += 1
    return ImportResult(added, updated, replaced, seen)


async def retire_missing(db: AsyncSession, source: str, type_: str, seen_ids: set[str]) -> int:
    """同源同类型里这次没出现的条目标记 retired（GEM 会把退役/取消的移出 operating）。"""
    rows = (
        (
            await db.execute(
                select(CatalogPlant.id).where(
                    CatalogPlant.source == source,
                    CatalogPlant.type == type_,
                    CatalogPlant.status == "operating",
                )
            )
        )
        .scalars()
        .all()
    )
    gone = [i for i in rows if i not in seen_ids]
    if gone:
        await db.execute(
            update(CatalogPlant).where(CatalogPlant.id.in_(gone)).values(status="retired")
        )
    return len(gone)
