"""目录风电的机型档与海上格点。docs/07 §2.2、docs/04 §二 / §七

依据是 REIT 场站电量对账（docs/07 §8.1 2026-09-15）：通用功率曲线把 2015 年后的陆上
低风速机型少算约一半；上游默认 land 格点把近岸海上风电分到陆地，电量少算三成。
"""

from datetime import datetime, timedelta
from zoneinfo import ZoneInfo

import httpx
import openpyxl
import pytest
import respx

from app.catalog import importer
from app.config import settings
from app.models import CatalogPlant
from app.services import fleet_prediction as fleet
from app.services import weather
from app.services.catalog import is_offshore
from app.services.prediction_basis import (
    catalog_cell_selection,
    catalog_turbine_class,
    wind_turbine_class,
)
from app.services.station import from_catalog
from tests.fixtures_forecast import TZ, make_forecast

MODERN = settings.wind_catalog_modern_from_year


def _plant(
    name="Gansu Guazhou wind farm",
    local=None,
    year=2020,
    phases=None,
    type_="wind",
    lat=40.52,
    lon=95.78,
):
    return CatalogPlant(
        id="gem:L1",
        source="gem",
        source_id="L1",
        type=type_,
        name=name,
        name_local=local,
        latitude=lat,
        longitude=lon,
        capacity_kw=100_000,
        commissioning_year=year,
        status="operating",
        provenance={"phases": phases if phases is not None else [{"id": "G1", "capacity_kw": 1e5}]},
    )


class TestOffshore:
    def test_旧库按名称识别海上风电(self):
        assert is_offshore(_plant(name="Jiangsu Binhai North Area H1 Offshore wind farm"))
        assert is_offshore(_plant(name="Binhai H2", local="国家电投江苏滨海北H2#400MW海上风电项目"))
        assert not is_offshore(_plant())
        assert not is_offshore(_plant(name="Offshore floating solar", type_="solar"))

    def test_分期安装类型优先于名称(self):
        onshore = [{"id": "G1", "capacity_kw": 1, "installation_type": "Onshore"}]
        offshore = [{"id": "G1", "capacity_kw": 1, "installation_type": "Offshore hard mount"}]
        assert not is_offshore(_plant(name="Haishang Offshore Road wind farm", phases=onshore))
        assert is_offshore(_plant(phases=offshore))


class TestTurbineClass:
    def test_陆上按投运年份_年份未知按新站_海上通用曲线(self):
        assert wind_turbine_class(MODERN, offshore=False) == settings.wind_catalog_modern_class
        assert wind_turbine_class(None, offshore=False) == settings.wind_catalog_modern_class
        assert wind_turbine_class(MODERN - 1, offshore=False) is None
        assert wind_turbine_class(MODERN + 5, offshore=True) is None

    def test_多期场址按容量加权的中位年份(self):
        phases = [
            {"id": "G1", "capacity_kw": 30_000, "start_year": MODERN - 4},
            {"id": "G2", "capacity_kw": 70_000, "start_year": MODERN + 3},
        ]
        # 目录记录的年份取自首期，只看它会把以新机型为主的场址判成老机型
        modern = _plant(year=MODERN - 4, phases=phases)
        assert catalog_turbine_class(modern) == settings.wind_catalog_modern_class
        old = [{**phases[0], "capacity_kw": 70_000}, {**phases[1], "capacity_kw": 30_000}]
        assert catalog_turbine_class(_plant(year=MODERN - 4, phases=old)) is None
        # 分期缺年份（导入该字段前的旧库）退回目录年份
        assert catalog_turbine_class(_plant(year=MODERN - 4)) is None

    def test_光伏没有机型档与格点选择(self):
        solar = _plant(type_="solar", name="Offshore PV")
        assert catalog_turbine_class(solar) is None
        assert catalog_cell_selection(solar) is None

    def test_公开电站带上机型档与海上格点(self):
        sea = from_catalog(_plant(name="Liaoning Dalian Zhuanghe 3 Offshore wind farm"))
        assert sea.turbine_class is None and sea._cell_selection == "sea"
        land = from_catalog(_plant())
        assert land.turbine_class == settings.wind_catalog_modern_class
        assert land._cell_selection is None


class TestImporter:
    def test_风电分期记录安装类型与投运年份(self, tmp_path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(
            [
                "Country/Area",
                "Project Name",
                "Phase Name",
                "Capacity (MW)",
                "Installation Type",
                "Status",
                "Start year",
                "Latitude",
                "Longitude",
                "GEM location ID",
                "GEM phase ID",
            ]
        )
        ws.append(
            [
                "China",
                "Binhai North H2 Offshore wind farm",
                "--",
                400,
                "Offshore hard mount",
                "operating",
                2018,
                34.36,
                120.21,
                "L1",
                "G1",
            ]
        )
        path = tmp_path / "Global-Wind-Power-Tracker.xlsx"
        wb.save(path)
        (row,) = importer.read_gem(path, "CHN")
        phase = row.provenance["phases"][0]
        assert phase["installation_type"] == "Offshore hard mount"
        assert phase["start_year"] == 2018


class TestPointForecast:
    @pytest.fixture(autouse=True)
    def _fresh(self):
        weather.clear_cache()
        yield
        weather.clear_cache()

    async def test_海上格点单独请求_单独缓存(self):
        now = datetime.now(ZoneInfo(TZ))
        start = (now - timedelta(days=1)).replace(
            hour=0, minute=0, second=0, microsecond=0, tzinfo=None
        )
        body = {"timezone": TZ, "minutely_15": make_forecast(start_date=start)["minutely_15"]}
        seen = []

        def respond(request: httpx.Request) -> httpx.Response:
            seen.append(request.url.params.get("cell_selection"))
            return httpx.Response(200, json=body)

        sea = from_catalog(_plant(name="Binhai Offshore wind farm"))
        land = from_catalog(_plant())  # 同一坐标
        with respx.mock(assert_all_called=False) as mock:
            mock.get(url__regex=r".*/static/meta\.json").respond(404)
            mock.get(url__regex=r".*/v1/forecast.*").mock(side_effect=respond)
            async with httpx.AsyncClient() as http:
                await weather.station_forecast(http, sea)
                await weather.station_forecast(http, land)
                await weather.station_forecast(http, sea)
        assert seen == ["sea", None]


class TestFleetCells:
    def test_海上风电与同格心陆上场站分开取气象(self):
        sea = fleet.cell(_plant(name="Binhai Offshore wind farm", lat=34.36, lon=120.21))
        land = fleet.cell(_plant(lat=34.36, lon=120.21))
        assert sea[:3] == land[:3]
        assert (sea[3], land[3]) == ("sea", None)
        assert fleet.coord_key(*sea[1:]) != fleet.coord_key(*land[1:])
        # 陆上键不变，已落盘的网格气象缓存继续命中
        assert fleet.coord_key(*land[1:]) == f"{land[1]},{land[2]}"

    async def test_全目录海上格点单独一次请求(self, tmp_path, monkeypatch):
        from tests.test_prediction import forecast, plant

        monkeypatch.setattr(fleet, "directory", lambda: tmp_path)
        raw = forecast()
        seen = []

        def respond(request: httpx.Request) -> httpx.Response:
            n = len(request.url.params["latitude"].split(","))
            seen.append((request.url.params.get("cell_selection"), n))
            return httpx.Response(200, json=[raw] * n if n > 1 else raw)

        # 同一格心：陆上与海上各一座
        plants = [plant("land"), plant("sea", name="Jiangsu Offshore wind farm")]
        with respx.mock(assert_all_called=False) as mock:
            mock.get(url__regex=r".*/static/meta\.json").respond(404)
            mock.get(url__regex=r".*/v1/forecast.*").mock(side_effect=respond)
            async with httpx.AsyncClient() as http:
                await fleet.build(http, "gfs_global", fleet.day_key(), plants)
        assert sorted(seen, key=str) == sorted([(None, 1), ("sea", 1)], key=str)
        out = fleet.load(tmp_path / f"gfs_global-{fleet.day_key()}.json")
        assert out["covered_count"] == 2
