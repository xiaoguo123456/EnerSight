"""导入公开电站目录。docs/04 §七、docs/06 §5.5

支持两个来源：
  wri  WRI Global Power Plant Database（CC BY 4.0，CSV 可直接下载）
       https://raw.githubusercontent.com/wri/global-power-plant-database/master/output_database/global_power_plant_database.csv
  gem  Global Energy Monitor 的 Global Solar / Wind Power Tracker（CC BY 4.0，官网填表下载 xlsx）
       中国覆盖最全且带中文名，导入后同址的 WRI 条目会被 GEM 覆盖

用法：
  uv run python scripts/import_catalog.py wri /path/to/global_power_plant_database.csv
  uv run python scripts/import_catalog.py gem /path/to/Global-Solar-Power-Tracker.xlsx
  uv run python scripts/import_catalog.py gem /path/to/Global-Wind-Power-Tracker.xlsx
  加 --country CHN 只导中国（默认）；--all 导全部国家

去重：同类型、相距 < 1.5 km 视为同一座，GEM 优先；同源按 source_id 幂等更新。
"""

from __future__ import annotations

import argparse
import asyncio
import csv
import re
import sys
from dataclasses import dataclass
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from sqlalchemy import select  # noqa: E402

from app.db import SessionLocal  # noqa: E402
from app.models import CatalogPlant  # noqa: E402
from app.services.catalog import haversine_km  # noqa: E402

DEDUP_KM = 1.5
MIN_CAPACITY_MW = 1.0
_JUNK = re.compile(r"[{}\[\]|<>]")  # WRI 源里偶有 "Qili}" 这种脏字符


def clean_name(s: str) -> str:
    return re.sub(r"\s+", " ", _JUNK.sub("", s)).strip()


@dataclass
class Row:
    source: str
    source_id: str
    name: str
    name_local: str | None
    type: str  # solar | wind
    capacity_mw: float
    lat: float
    lon: float
    province: str | None = None
    city: str | None = None
    district: str | None = None
    owner: str | None = None
    year: int | None = None
    status: str = "operating"

    @property
    def id(self) -> str:
        return f"{self.source}:{self.source_id}"[:24]


# ────────────────────────────── WRI ──────────────────────────────

_WRI_FUEL = {"Solar": "solar", "Wind": "wind"}


def read_wri(path: Path, country: str | None) -> list[Row]:
    out: list[Row] = []
    with path.open(encoding="utf-8") as f:
        for r in csv.DictReader(f):
            if country and r["country"] != country:
                continue
            t = _WRI_FUEL.get(r["primary_fuel"])
            if not t or not r["capacity_mw"] or not r["latitude"]:
                continue
            cap = float(r["capacity_mw"])
            if cap < MIN_CAPACITY_MW:
                continue
            year = r.get("commissioning_year") or None
            out.append(
                Row(
                    source="wri",
                    source_id=r["gppd_idnr"],
                    name=clean_name(r["name"]),
                    name_local=None,
                    type=t,
                    capacity_mw=cap,
                    lat=float(r["latitude"]),
                    lon=float(r["longitude"]),
                    owner=(r.get("owner") or "").strip() or None,
                    year=int(float(year)) if year else None,
                )
            )
    return out


# ────────────────────────────── GEM ──────────────────────────────


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]", "", str(h).lower())


def _col(headers: list[str], *candidates: str) -> int | None:
    """按规范化后的列名找列；候选按优先级，支持前缀匹配。"""
    normed = [_norm(h) for h in headers]
    for c in candidates:
        cn = _norm(c)
        for i, h in enumerate(normed):
            if h == cn:
                return i
    for c in candidates:
        cn = _norm(c)
        for i, h in enumerate(normed):
            if h.startswith(cn):
                return i
    return None


def read_gem(path: Path, country: str | None) -> list[Row]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # 数据表通常叫 Data；找第一张含 Latitude 列的表
    ws = None
    for sheet in wb.worksheets:
        head = [c for c in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        if _col([str(h) for h in head if h is not None], "Latitude") is not None:
            ws = sheet
            break
    if ws is None:
        raise SystemExit("没找到含 Latitude 列的数据表")
    rows = ws.iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows)]
    ix = {
        "country": _col(headers, "Country/Area", "Country"),
        "name": _col(headers, "Project Name", "Plant Name", "Name"),
        "local": _col(
            headers, "Project Name in Local Language / Script", "Local Name", "Chinese Name"
        ),
        "cap": _col(headers, "Capacity (MW)", "Capacity"),
        "status": _col(headers, "Status"),
        "lat": _col(headers, "Latitude"),
        "lon": _col(headers, "Longitude"),
        "owner": _col(headers, "Owner", "Operator"),
        "year": _col(headers, "Start year", "Start Year", "Year"),
        "prov": _col(headers, "State/Province", "Subnational unit (province, state)", "State"),
        "city": _col(headers, "Major area (prefecture, district)", "City"),
        "district": _col(headers, "Local area (taluk, county)", "Local area"),
        "loc_id": _col(headers, "GEM location ID", "GEM location"),
        "phase_id": _col(headers, "GEM phase ID", "GEM unit/phase ID", "GEM unit ID"),
    }
    missing = [k for k in ("name", "cap", "lat", "lon", "status") if ix[k] is None]
    if missing:
        raise SystemExit(f"GEM 表缺少列 {missing}，表头：{headers}")
    t = "solar" if "solar" in path.name.lower() else "wind" if "wind" in path.name.lower() else None
    tech = _col(headers, "Technology Type", "Type")
    country_alias = {"CHN": "china"}

    def cell(r, k):
        i = ix[k]
        return r[i] if i is not None and i < len(r) else None

    grouped: dict[str, Row] = {}
    for r in rows:
        if country:
            c = str(cell(r, "country") or "").lower()
            if country_alias.get(country, country.lower()) not in c:
                continue
        status = str(cell(r, "status") or "").lower()
        if status not in ("operating",):
            continue
        cap, lat, lon = cell(r, "cap"), cell(r, "lat"), cell(r, "lon")
        if cap is None or lat is None or lon is None:
            continue
        try:
            cap, lat, lon = float(cap), float(lat), float(lon)
        except (TypeError, ValueError):
            continue
        if cap < MIN_CAPACITY_MW:
            continue
        kind = t
        if kind is None and tech is not None:
            kind = "wind" if "wind" in str(r[tech] or "").lower() else "solar"
        if kind is None:
            continue
        loc = str(cell(r, "loc_id") or cell(r, "phase_id") or f"{lat:.4f},{lon:.4f}")
        year = cell(r, "year")
        try:
            year = int(float(year)) if year not in (None, "") else None
        except (TypeError, ValueError):
            year = None
        key = f"{kind}:{loc}"
        if key in grouped:
            grouped[key].capacity_mw += cap  # 同一场址多期合并
            continue
        grouped[key] = Row(
            source="gem",
            source_id=loc,
            name=clean_name(str(cell(r, "name") or "")),
            name_local=(clean_name(str(cell(r, "local"))) or None) if cell(r, "local") else None,
            type=kind,
            capacity_mw=cap,
            lat=lat,
            lon=lon,
            province=(str(cell(r, "prov")).strip() or None) if cell(r, "prov") else None,
            city=(str(cell(r, "city")).strip() or None) if cell(r, "city") else None,
            district=(str(cell(r, "district")).strip() or None) if cell(r, "district") else None,
            owner=(str(cell(r, "owner")).strip() or None) if cell(r, "owner") else None,
            year=year,
        )
    return list(grouped.values())


# ────────────────────────────── 写库 ──────────────────────────────


async def upsert(rows: list[Row]) -> tuple[int, int, int]:
    """返回 (新增, 更新, 因同址跳过/替换)。"""
    added = updated = replaced = 0
    async with SessionLocal() as db:
        existing = (await db.execute(select(CatalogPlant))).scalars().all()
        by_id = {p.id: p for p in existing}
        # 粗网格索引做同址查找
        grid: dict[tuple[str, int, int], list[CatalogPlant]] = {}
        for p in existing:
            grid.setdefault((p.type, int(p.latitude * 20), int(p.longitude * 20)), []).append(p)

        def near_dupes(r: Row) -> list[CatalogPlant]:
            out = []
            gy, gx = int(r.lat * 20), int(r.lon * 20)
            for dy in (-1, 0, 1):
                for dx in (-1, 0, 1):
                    for p in grid.get((r.type, gy + dy, gx + dx), []):
                        if (
                            p.id != r.id
                            and haversine_km(r.lat, r.lon, p.latitude, p.longitude) < DEDUP_KM
                        ):
                            out.append(p)
            return out

        for r in rows:
            dupes = near_dupes(r)
            if dupes:
                if r.source == "wri" and any(p.source == "gem" for p in dupes):
                    replaced += 1
                    continue  # GEM 已有更好的记录
                if r.source == "gem":
                    for p in dupes:
                        if p.source == "wri":
                            await db.delete(p)
                            replaced += 1
            p = by_id.get(r.id)
            fields = dict(
                source=r.source,
                source_id=r.source_id,
                name=r.name[:128],
                name_local=(r.name_local or None) and r.name_local[:128],
                type=r.type,
                capacity_kw=r.capacity_mw * 1000.0,
                latitude=r.lat,
                longitude=r.lon,
                province=r.province,
                city=r.city,
                district=r.district,
                owner_name=(r.owner or None) and r.owner[:128],
                commissioning_year=r.year,
                status=r.status,
            )
            if p:
                for k, v in fields.items():
                    setattr(p, k, v)
                updated += 1
            else:
                p = CatalogPlant(id=r.id, **fields)
                db.add(p)
                by_id[r.id] = p
                grid.setdefault((r.type, int(r.lat * 20), int(r.lon * 20)), []).append(p)
                added += 1
        await db.commit()
    return added, updated, replaced


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("source", choices=["wri", "gem"])
    ap.add_argument("path", type=Path)
    ap.add_argument("--country", default="CHN", help="ISO3，默认 CHN；--all 导全部")
    ap.add_argument("--all", action="store_true")
    args = ap.parse_args()
    country = None if args.all else args.country
    rows = read_wri(args.path, country) if args.source == "wri" else read_gem(args.path, country)
    n_solar = sum(1 for r in rows if r.type == "solar")
    print(f"读到 {len(rows)} 条（光伏 {n_solar}，风电 {len(rows) - n_solar}），写库 ...")
    added, updated, replaced = asyncio.run(upsert(rows))
    print(f"新增 {added}，更新 {updated}，同址去重 {replaced}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
